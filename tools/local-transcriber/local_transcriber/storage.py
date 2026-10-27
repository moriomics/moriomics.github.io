from __future__ import annotations
from contextlib import contextmanager
from dataclasses import asdict
from datetime import datetime
import json
import os
from pathlib import Path
import tempfile
from .domain import Project, Settings, now


def app_root():
    return Path(os.environ.get('LOCALAPPDATA', Path.home() / '.local' / 'share')) / 'LocalTranscriber'


def app_home():
    root = app_root()
    root.mkdir(parents=True, exist_ok=True)
    return root


def settings_path():
    # Diagnostics must work even when creating the directory fails.
    return app_root() / 'settings.json'


def storage_error_message(exc):
    import errno
    code = getattr(exc, 'errno', None)
    win = getattr(exc, 'winerror', None)
    if win in (32, 33):
        reason = '別のアプリやプロセスがファイルを使用しています。'
    elif win == 5 or code in (errno.EACCES, errno.EPERM):
        reason = '書き込みが拒否されました。保存先の権限・読み取り専用属性・保護ソフトの制限などを確認してください。'
    elif win in (39, 112) or code == errno.ENOSPC:
        reason = '保存先の空き容量が不足しています。'
    elif win in (2, 3) or code == errno.ENOENT:
        reason = '保存先のフォルダまたはファイルが見つかりません。'
    elif win == 6 or code == errno.EBADF:
        reason = '保存処理で使うファイルハンドルが無効です。'
    elif isinstance(exc, (ValueError, TypeError)):
        reason = '保存するデータまたは設定の形式が不正です。'
    else:
        reason = '原因は特定できません。下のエラー情報を確認してください。'
    # Never include exception text: it can contain user content.
    stage = getattr(exc, 'storage_stage', '設定の検証・保存')
    return (f'原因: {reason}\n工程: {stage}\n'
            f'エラー種別: {type(exc).__name__} / errno={code} / winerror={win}')


@contextmanager
def atomic_target(path):
    path = Path(path)
    tmp = None
    stage = '保存先の準備'
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        fd, name = tempfile.mkstemp(prefix='.' + path.stem + '-', suffix=path.suffix, dir=path.parent)
        os.close(fd)
        tmp = Path(name)
        stage = '一時ファイルの書き込み'
        yield tmp
        stage = 'ディスクへの確定'
        # Windows flushing requires a writable handle. rb can fail at fsync.
        with tmp.open('r+b') as f:
            f.flush()
            os.fsync(f.fileno())
        stage = '保存ファイルの置換'
        os.replace(tmp, path)
    except OSError as exc:
        exc.storage_stage = stage
        raise
    finally:
        if tmp is not None:
            try:
                tmp.unlink(missing_ok=True)
            except OSError:
                # Preserve the original error if Windows still has a lock.
                pass


def write_json(path, data):
    with atomic_target(path) as tmp:
        tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')


def protected_target(project, path):
    if Path(path).resolve() == Path(project.source_path).resolve():
        raise ValueError('元の音声・動画ファイルには保存できません。')


def save_project(project, path):
    protected_target(project, path)
    if Path(path).suffix.lower() != '.json':
        raise ValueError('プロジェクトの拡張子は.jsonにしてください。')
    project.updated_at = now()
    write_json(path, project.to_dict())


def load_project(path):
    path = Path(path)
    if path.stat().st_size > 100 * 1024 * 1024:
        raise ValueError('100 MiBを超えるプロジェクトはMVPでは読み込めません。')
    return Project.from_dict(json.loads(path.read_text(encoding='utf-8')))


def save_settings(settings):
    settings.validate()
    write_json(settings_path(), asdict(settings))


def load_settings():
    try:
        settings = Settings(**json.loads((app_home() / 'settings.json').read_text(encoding='utf-8')))
        settings.validate()
        return settings
    except (OSError, ValueError, TypeError):
        return Settings()


def timestamp(seconds, srt=False):
    ms = max(0, round(seconds * 1000))
    hours, ms = divmod(ms, 3600000)
    minutes, ms = divmod(ms, 60000)
    sec, ms = divmod(ms, 1000)
    return f'{hours:02}:{minutes:02}:{sec:02}{"," if srt else "."}{ms:03}'


def export_project(project, path, timestamps=True):
    protected_target(project, path)
    path = Path(path)
    ext = path.suffix.lower()
    segs = sorted(project.segments, key=lambda s: (s.start, s.end, s.id))
    for s in segs:
        s.validate()
    if ext == '.json':
        save_project(project, path)
        return
    with atomic_target(path) as tmp:
        if ext == '.txt':
            lines = [f'[{timestamp(s.start)} - {timestamp(s.end)}] {s.text}' if timestamps else s.text for s in segs]
            tmp.write_text('\n'.join(lines), encoding='utf-8')
        elif ext == '.srt':
            # Preserve authored timestamps (including legitimate overlaps).
            blocks = [f'{i}\n{timestamp(s.start, True)} --> {timestamp(s.end, True)}\n{s.text.replace(chr(13), "").replace(chr(10)+chr(10), chr(10))}\n'
                      for i, s in enumerate((x for x in segs if x.text.strip()), 1)]
            tmp.write_text('\n'.join(blocks), encoding='utf-8')
        elif ext == '.docx':
            from docx import Document
            doc = Document()
            doc.add_heading('文字起こし', 0)
            doc.add_paragraph('ファイル名: ' + project.filename)
            doc.add_paragraph('作成日: ' + datetime.now().astimezone().isoformat(timespec='seconds'))
            for s in segs:
                if timestamps:
                    doc.add_paragraph(f'{timestamp(s.start)} - {timestamp(s.end)}', style='Caption')
                doc.add_paragraph(s.text)
            doc.save(tmp)
        else:
            raise ValueError('出力形式はTXT / SRT / JSON / DOCXを指定してください。')
